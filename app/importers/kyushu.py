"""Validated adapter for the supplied Kyushu 2026-10 code-content listing.

Unknown layouts fail closed. Blank bed cells remain unknown, never fabricated zero.
All original nonempty rows are retained for reconciliation.
"""
import hashlib
import re
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
import xlrd
from app.domain import PREFECTURES, REGIONS, BED_LABELS, classify, institution_code, normalize

VERSION = 'code-content-v3'
MAX_SIZE = 40 * 1024 * 1024
DATES = re.compile(r'(令和|平成|昭和|令|平|昭)\s*(元|\d+)\s*[年.]\s*(\d+)\s*[月.]\s*(\d+)\s*日?')

def japanese_date(value):
    text = normalize(value)
    m = DATES.search(text)
    if m:
        era, y, month, day = m.groups()
        offset = 2018 if era.startswith('令') else 1988 if era.startswith('平') else 1925
        return date(offset + (1 if y == '元' else int(y)), int(month), int(day)).isoformat()
    m = re.search(r'(20\d{2})[年/.-]\s*(\d+)[月/.-]\s*(\d+)',text)
    return date(*map(int,m.groups())).isoformat() if m else None

def workbook_rows(blob, filename):
    if filename.lower().endswith('.xlsx'):
        with ZipFile(BytesIO(blob)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 150 * 1024 * 1024:
                raise ValueError('Excel展開サイズ超過')
        book = load_workbook(BytesIO(blob), read_only=True, data_only=True)
        try:
            result=[]
            for s in book:
                # Some official workbooks incorrectly advertise A1:A1 as their dimension.
                if s.max_row==1 and s.max_column==1: s.reset_dimensions()
                result.append((s.title,list(s.values)))
            return result
        finally:
            book.close()
    book = xlrd.open_workbook(file_contents=blob)
    return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]

def owner_group(text):
    for marker, label in [('医療法人','医療法人'),('国立病院機構','独立行政法人'),
                          ('独立行政法人','独立行政法人'),('社会福祉法人','社会福祉法人'),
                          ('財団法人','公益・一般法人'),('社団法人','公益・一般法人'),
                          ('日本赤十字','日本赤十字社'),('学校法人','学校法人'),
                          ('株式会社','会社')]:
        if marker in text: return label
    return '未分類'  # Do not assume individual ownership from a missing corporate marker.

def parse_block(block, pref, source):
    first = block[0][1]
    text = lambda col: normalize(first[col])
    bed_values = {}
    pending = None
    ambiguous_beds = False
    departments = []
    raw_bed_text = []
    labels = {'一般':'general','療養':'medical','精神':'psychiatric','結核':'tuberculosis',
              '感染':'infectious','感染症':'infectious'}
    for _, row in block:
        value = normalize(row[8])
        if not value: continue
        raw_bed_text.append(value)
        # A wrapped label e.g. 一般(感染) + "4" must take precedence over 一般.
        match = re.fullmatch(r'(一般\(感染\)|一般|療養|精神|結核|感染症?)(?:\s+([\d,]+))?',value)
        if match:
            label, number = match.groups()
            key = 'infectious' if '感染' in label else labels[label]
            pending = key
            if number is None: continue
        elif pending and re.fullmatch(r'[\d,]+',value):
            key, number = pending, value
        else:
            if pending: ambiguous_beds = True
            departments.append(value)
            pending = None
            continue
        n = int(number.replace(',',''))
        # Separate subcategories can be printed on multiple rows of the same bed type.
        bed_values[key] = bed_values.get(key,0)+n
        pending = None
    if pending: ambiguous_beds = True
    total = sum(bed_values.values()) if bed_values and not ambiguous_beds else None
    # Remarks such as 地域支援 can precede a wrapped 病院 label.
    kinds={('病院' if normalize(row[9])=='総合病院' else normalize(row[9]))
           for _,row in block if normalize(row[9]) in ['病院','総合病院','診療所']}
    explicit = next(iter(kinds)) if len(kinds)==1 else None
    code = institution_code(text(1),pref)
    override = None
    if code == '1112802007':
        if text(2) != '自衛隊入間病院' or explicit != '病院' or total != 2:
            raise ValueError('個別区分判定の前提が変更されました。再確認が必要です')
        override = {'kind':'診療所','basis':'2026-10-15 ユーザー指定：法律上は診療所として扱う','source_kind':explicit}
    classification = classify(total, override['kind'] if override else explicit, bed_values.get('psychiatric'))
    classification_conflict=classification['needs_review'] or len(kinds)>1
    if classification_conflict: classification['needs_review']=True
    # Absence of a psychiatric row alone is not proof of zero psychiatric beds.
    if ambiguous_beds: classification['needs_review'] = True
    statuses = {normalize(row[9]) for _,row in block if normalize(row[9]) in ['現存','休止','廃止','辞退','取消']}
    status = next(iter(statuses)) if len(statuses)==1 else '確認中'
    reasons = [normalize(row[7]) for _,row in block[1:] if normalize(row[7]) and not japanese_date(row[7])]
    renewal = [japanese_date(row[7]) for _,row in block[1:] if japanese_date(row[7])]
    address = text(3)
    postal = re.search(r'〒?\s*(\d{3})[-ー−]\s*(\d{4})',address)
    postal_code = '-'.join(postal.groups()) if postal else None
    address = re.sub(r'^〒?\s*\d{3}[-ー−]\s*\d{4}', '',address).strip()
    local = address.removeprefix(PREFECTURES[pref])
    city = re.match(r'(.+?市.+?区|.+?市|.+?郡.+?[町村]|.+?区|.+?[町村])',local)
    return {
        'code':institution_code(text(1),pref), 'prefecture_code':pref, 'prefecture':PREFECTURES[pref],
        'bureau':source['bureau'], 'municipality':city.group(1) if city else None,
        'name':text(2), 'postal_code':postal_code,'address':address,'phone':text(4),
        'owner':text(5),'owner_type':owner_group(text(5)),'manager':text(6),
        'departments':departments,'total_beds':total,
        **{f'{key}_beds':bed_values.get(key) for key in BED_LABELS},
        # Unknown bed notation is a data-quality flag, not a facility closure or
        # unknown operating status. Numeric kind/bed conflicts remain held.
        **classification,'status':'確認中' if classification_conflict else status,
        'designation_date':japanese_date(first[7]),'renewal_date':renewal[-1] if renewal else None,
        'change_reason':' / '.join(reasons) or None,'opening_classification':'判定保留',
        'reference_date':source['reference_date'],'file_created_date':source['file_created_date'],
        'retrieved_at':source['retrieved_at'],'source_file':source['filename'],
        'raw':{'sheet':source['sheet'],'start_row':block[0][0],
               'rows':[{'row':i,'cells':list(row[:10])} for i,row in block],
               'bed_interpretation':'Blank categories remain NULL; unresolved bed labels make total NULL',
               'explicit_kind':explicit,'source_status':status,
               'classification_conflict':classification_conflict,'bed_text':raw_bed_text,
               **({'classification_override':override} if override else {})},
    }

def parse_workbook(blob, filename, pref, retrieved_at=None, *, bureau=None, allow_missing_prefecture=False, _sheets=None):
    bureau=bureau or next(b for b,ps in REGIONS.items() if pref in ps)
    if pref not in REGIONS[bureau]: raise ValueError('厚生局と担当県が一致しません')
    results, sources = [], []
    for sheet, original_rows in (_sheets if _sheets is not None else workbook_rows(blob,filename)):
        rows = [tuple(normalize(v) for v in row) for row in original_rows]
        header = '\n'.join(' '.join(r) for r in rows[:11])
        # Content proof, independent of filename; reject dental / pharmacy listings.
        if not re.search(r'現在\s+医科\s',header):
            if '現在' in header and ('歯科' in header or '薬局' in header): continue
            raise ValueError(f'未対応の表形式: {filename}/{sheet}')
        if (PREFECTURES[pref] not in header and not allow_missing_prefecture) or 'コード内容別医療機関一覧表' not in header:
            raise ValueError('都道府県または表の種類が一致しません')
        if any(len(r)>10 and any(r[10:]) for r in rows):
            raise ValueError('列構造が変わりました。検証が必要です')
        ref = [japanese_date(v) for row in rows[:11] for v in row if '現在' in v]
        created = [japanese_date(v) for row in rows[:11] for v in row if '作成' in v]
        if not ref or not ref[0] or not created or not created[0]: raise ValueError('元資料の日付が不明です')
        source = {'filename':filename,'sheet':sheet,'prefecture_code':pref,'bureau':bureau,
                  'reference_date':ref[0],'file_created_date':created[0],
                  'retrieved_at':retrieved_at or datetime.now(timezone.utc).isoformat(),
                  'sha256':hashlib.sha256(blob).hexdigest(),'parser_version':VERSION}
        blocks, block, sequence = [], [], []
        for i,row in enumerate(rows,1):
            row = row + ('',)*max(0,10-len(row))
            if re.fullmatch(r'\d+',row[0]) and row[1] and row[2]:
                if block: blocks.append(block)
                institution_code(row[1],pref)
                sequence.append(int(row[0]))
                block = [(i,row)]
            elif block:
                if row[0] or (row[1] and re.fullmatch(r'[\d,・. -]+',row[1])):
                    raise ValueError(f'未対応の補足行 {i}: {row[:2]}')
                block.append((i,row))
        if block: blocks.append(block)
        if not sequence or sequence != list(range(1,len(sequence)+1)):
            raise ValueError('施設連番の欠落・重複・空ファイルを検出しました')
        parsed = [parse_block(b,pref,source) for b in blocks]
        if len(set(r['code'] for r in parsed)) != len(parsed): raise ValueError('医科施設コード重複')
        if any(not r['name'] or not r['address'] for r in parsed): raise ValueError('必須情報欠落')
        if any(r['raw'].get('classification_override') for r in parsed):
            source['parser_version'] = VERSION + '-kind1'
        source['record_count'] = len(parsed)
        results.extend(parsed); sources.append(source)
    return results,sources

def parse_archive(path, pref, retrieved_at=None):
    path = Path(path)
    if path.stat().st_size > MAX_SIZE: raise ValueError('ZIPサイズ超過')
    records, sources = [], []
    with ZipFile(path) as z:
        if len(z.infolist())>100 or sum(i.file_size for i in z.infolist())>MAX_SIZE:
            raise ValueError('ZIP展開サイズ超過')
        for entry in z.infolist():
            if not entry.filename.lower().endswith(('.xls','.xlsx')): continue
            # User scope: inspect medicine and dentistry; pharmacy content is unnecessary.
            if 'yakkyoku' in entry.filename.lower(): continue
            # Inspect dental content too; never select medical data by filename alone.
            data, meta = parse_workbook(z.read(entry),entry.filename,pref,retrieved_at)
            records.extend(data); sources.extend(meta)
    if not records: raise ValueError('医科データがありません')
    if len({r['code'] for r in records}) != len(records): raise ValueError('ZIP内コード重複')
    for src in sources:
        src['archive_sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        src['archive_name']=path.name
    return records,sources
