from io import BytesIO
from zipfile import ZipFile
import pytest
from openpyxl import Workbook,load_workbook
from app.importers.kyushu import parse_workbook,parse_archive

def blob(category='医科',bed_rows=None,code='021,003,9'):
    book=Workbook();s=book.active
    s.cell(3,1,'コード内容別医療機関一覧表');s.cell(3,2,'[福岡県]')
    s.cell(5,1,f'[令和 8年10月 1日現在　{category}　現存/休止]');s.cell(5,2,'令和8年10月5日作成')
    first=['1',code,'歯科併設医科病院','〒812－0063福岡県福岡市東区原田','092-000-0000','医療法人テスト','管理者','平14.6.1', '一般 80','病院']
    for i,v in enumerate(first,1):s.cell(12,i,v)
    continuation=['','(033,864,6)','','','','','','組織変更','精神 20','現存']
    for i,v in enumerate(continuation,1):s.cell(13,i,v)
    for row in bed_rows or []:s.append(row)
    output=BytesIO();book.save(output);return output.getvalue()

def test_multiline_beds_and_supplemental_dental_code():
    records,sources=parse_workbook(blob(),'anything.xlsx',40)
    assert len(records)==1
    r=records[0]
    assert r['code']=='4010210039' and r['total_beds']==100
    assert r['general_beds']==80 and r['psychiatric_beds']==20
    assert r['medical_beds'] is None
    assert r['change_reason']=='組織変更' and r['opening_classification']=='判定保留'
    assert r['reference_date']=='2026-10-01' and r['file_created_date']=='2026-10-05'

def test_dental_content_excluded_even_if_filename_says_ika():
    assert parse_workbook(blob('歯科'),'ika.xlsx',40)==([],[])

def test_wrong_prefecture_rejected():
    with pytest.raises(ValueError): parse_workbook(blob(),'f.xlsx',41)

def test_unknown_layout_fails_closed():
    b=Workbook();out=BytesIO();b.save(out)
    with pytest.raises(ValueError):parse_workbook(out.getvalue(),'new-layout.xlsx',40)

def test_pharmacy_is_not_parsed(tmp_path):
    path=tmp_path/'source.zip'
    with ZipFile(path,'w') as z:
        z.writestr('ika.xlsx',blob());z.writestr('yakkyoku.xlsx',b'not-an-excel-file')
    records,_=parse_archive(path,40)
    assert len(records)==1

def test_medical_counterpart_annex_is_reported_without_prefixing_dental_code(tmp_path):
    from app.importers.national import parse_region
    path=tmp_path/'medical.zip'
    with ZipFile(path,'w') as z:
        z.writestr('main.xlsx',blob())
        z.writestr('annex.xlsx',blob('医科併設',code='033,864,6'))
    exclusions=[]
    records,sources,errors=parse_region(path,'九州厚生局',excluded_sources=exclusions)
    assert errors==[]
    assert [r['code'] for r in records]==['4010210039']
    assert len(sources)==1
    assert len(exclusions)==1
    assert exclusions[0]['publication_category']=='医科併設'
    assert exclusions[0]['primary_rows']==1
    assert exclusions[0]['filename']=='annex.xlsx'

@pytest.mark.parametrize('status',['現存','休止'])
def test_unknown_bed_label_keeps_explicit_facility_status(status):
    book=load_workbook(BytesIO(blob()))
    sheet=book.active
    sheet.cell(12,9,'一般')
    sheet.cell(12,10,'診療所')
    sheet.cell(13,9,'内')
    sheet.cell(13,10,status)
    out=BytesIO();book.save(out)
    records,_=parse_workbook(out.getvalue(),'main.xlsx',40)
    row=records[0]
    assert row['kind']=='診療所' and row['total_beds'] is None
    assert row['needs_review'] is True and row['status']==status
    assert row['raw']['source_status']==status

def test_wrapped_hospital_kind_is_read_from_continuation_row():
    book=load_workbook(BytesIO(blob()))
    sheet=book.active
    sheet.cell(12,10,'地域支援')
    sheet.cell(13,10,'病院')
    sheet.append(['','','','','','','','','','現存'])
    out=BytesIO();book.save(out)
    records,_=parse_workbook(out.getvalue(),'main.xlsx',40)
    assert records[0]['kind']=='病院' and records[0]['status']=='現存'

def test_explicit_numeric_kind_conflict_remains_held():
    book=load_workbook(BytesIO(blob()))
    sheet=book.active
    sheet.cell(12,9,'一般 2')
    sheet.cell(13,9,'内')
    out=BytesIO();book.save(out)
    records,_=parse_workbook(out.getvalue(),'main.xlsx',40)
    assert records[0]['total_beds']==2
    assert records[0]['needs_review'] and records[0]['status']=='確認中'
