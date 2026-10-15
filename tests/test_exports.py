from io import BytesIO
from openpyxl import load_workbook
from app.exports import workbook_bytes,csv_bytes

def test_workbook_scopes_filters_formula_and_codes():
    rows=[{'code':'0110000001','name':'=HYPERLINK("bad")','status':'現存','kind':'病院','prefecture':'北海道','total_beds':100,'psychiatric_beds':100},
          {'code':'0110000002','name':'廃止施設','status':'廃止','kind':'病院'}]
    book=load_workbook(BytesIO(workbook_bytes(rows,[])),data_only=False)
    assert book.sheetnames==['検索・抽出フォーム','集計ダッシュボード','病院','診療所','精神科病院','取得元・更新情報']
    assert book['病院'].max_row==2 and book['精神科病院'].max_row==2
    assert book['病院']['A2'].value=='0110000001' and book['病院']['A2'].number_format=='@'
    assert book['病院']['B2'].value.startswith("'")
    assert book['病院'].auto_filter.ref=='A1:W2'
    assert '_xlfn._xlws.FILTER' in book['検索・抽出フォーム']['A13'].value
    assert len(book['検索・抽出フォーム'].data_validations.dataValidation)==3
    assert '廃止施設' not in str([[c.value for c in r] for r in book['病院']])

def test_csv_formula_sanitized_and_utf8_bom():
    data=csv_bytes([{'code':'0110000001','name':'=1+1'}])
    assert data.startswith(b'\xef\xbb\xbf') and "'=1+1" in data.decode('utf-8-sig')
